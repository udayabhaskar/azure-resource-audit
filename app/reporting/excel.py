from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)

from app.models.resource import AzureResource


def export_inventory(
    resources: list[AzureResource],
    subscription_name: str,
    subscription_id: str,
    output_directory: str = "output",
) -> Path:
    """
    Generate an Excel report containing the Azure
    resource inventory.
    """

    # --------------------------------------------------
    # Output directory
    # --------------------------------------------------

    output_path = Path(
        output_directory
    )

    output_path.mkdir(
        parents=True,
        exist_ok=True,
    )

    # --------------------------------------------------
    # File name
    # --------------------------------------------------

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    filename = (
        "Azure_Resource_Audit_"
        f"{timestamp}.xlsx"
    )

    file_path = (
        output_path / filename
    )

    # --------------------------------------------------
    # Workbook
    # --------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Resource Inventory"
    )

    # --------------------------------------------------
    # Headers
    # --------------------------------------------------

    headers = [
        "Resource ID",
        "RG",
        "Resource",
        "Type",
        "Location",
        "State",
        "SKU",
        "Monthly Cost",
        "Owner",
        "Environment",
        "Tags",
    ]

    worksheet.append(headers)

    # --------------------------------------------------
    # Resource data
    # --------------------------------------------------

    for resource in resources:

        monthly_cost = (
            resource.monthly_cost
            if resource.monthly_cost
            is not None
            else "N/A"
        )

        worksheet.append(
            [
                resource.resource_id,
                resource.resource_group,
                resource.resource_name,
                resource.resource_type,
                resource.location,
                resource.state,
                resource.sku,
                monthly_cost,
                resource.owner,
                resource.environment,
                resource.tags,
            ]
        )

    # --------------------------------------------------
    # Header formatting
    # --------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    header_alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    for cell in worksheet[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = (
            header_alignment
        )

    # --------------------------------------------------
    # Worksheet behavior
    # --------------------------------------------------

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # --------------------------------------------------
    # Column widths
    # --------------------------------------------------

    widths = {
        "A": 70,
        "B": 30,
        "C": 40,
        "D": 45,
        "E": 20,
        "F": 18,
        "G": 25,
        "H": 18,
        "I": 25,
        "J": 20,
        "K": 60,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------
    # Monthly Cost formatting
    # --------------------------------------------------

    for cell in worksheet["H"][1:]:

        if isinstance(
            cell.value,
            (int, float),
        ):

            cell.number_format = (
                '#,##0.00'
            )

    # --------------------------------------------------
    # Audit Information sheet
    # --------------------------------------------------

    metadata = workbook.create_sheet(
        "Audit Information"
    )

    metadata_data = [
        (
            "Subscription Name",
            subscription_name,
        ),
        (
            "Subscription ID",
            subscription_id,
        ),
        (
            "Generated At",
            datetime.now().isoformat(),
        ),
        (
            "Total Resources",
            len(resources),
        ),
    ]

    for row in metadata_data:

        metadata.append(row)

    metadata.column_dimensions[
        "A"
    ].width = 25

    metadata.column_dimensions[
        "B"
    ].width = 70

    for cell in metadata["A"]:

        cell.font = Font(
            bold=True
        )

    # --------------------------------------------------
    # Save workbook
    # --------------------------------------------------

    workbook.save(file_path)

    return file_path